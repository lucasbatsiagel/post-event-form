import io
import json
import os
import re
from datetime import datetime, timedelta, date
from functools import wraps

import qrcode
import qrcode.image.svg
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from dotenv import load_dotenv
from flask import Flask, render_template, request, redirect, url_for, session, flash, jsonify, abort, Response, send_file
from sqlalchemy import text

from models import db, Event, Submission

load_dotenv()

def _normalized_db_url():
    url = os.environ.get("DATABASE_URL", "sqlite:///pullsheets.db")
    # Render/Heroku-style URLs use the old "postgres://" scheme, and some
    # providers (e.g. Neon) hand out "postgresql+psycopg://" (driver v3).
    # Force the psycopg2 driver, which is the one actually installed.
    if url.startswith(("postgres://", "postgresql://", "postgresql+")):
        url = re.sub(r"^postgres(ql)?(\+\w+)?://", "postgresql+psycopg2://", url, count=1)
    return url


app = Flask(__name__)
app.config["SECRET_KEY"] = os.environ.get("SECRET_KEY", "dev-secret-change-me")
app.config["SQLALCHEMY_DATABASE_URI"] = _normalized_db_url()
app.config["SQLALCHEMY_TRACK_MODIFICATIONS"] = False
# Neon (and similar serverless Postgres) suspends its compute after a period
# of inactivity; the next connection from the pool may already be dead.
# pool_pre_ping tests each connection before use and transparently
# reconnects instead of raising, and pool_recycle avoids handing out
# connections the server may have already dropped.
app.config["SQLALCHEMY_ENGINE_OPTIONS"] = {
    "pool_pre_ping": True,
    "pool_recycle": 280,
}

ADMIN_PIN = os.environ.get("ADMIN_PIN", "1234")
REPORT_TOKEN = os.environ.get("REPORT_TOKEN", "")

db.init_app(app)

with app.app_context():
    db.create_all()
    # db.create_all() only creates tables that don't exist yet; it won't add
    # new columns to a table that's already there. Patch those in by hand.
    with db.engine.connect() as conn:
        if db.engine.dialect.name == "postgresql":
            conn.execute(text("ALTER TABLE event ADD COLUMN IF NOT EXISTS venue VARCHAR(200)"))
            conn.commit()
        else:
            existing_cols = [row[1] for row in conn.execute(text("PRAGMA table_info(event)"))]
            if "venue" not in existing_cols:
                conn.execute(text("ALTER TABLE event ADD COLUMN venue VARCHAR(200)"))
                conn.commit()


LOGO_PATH = os.path.join(app.root_path, "static", "img", "siagel-logo.png")

BRAND_RED = RGBColor(0xD3, 0x1F, 0x2B)
BRAND_INK = RGBColor(0x18, 0x18, 0x18)
BRAND_GRAY = RGBColor(0x58, 0x56, 0x5A)


def admin_required(view):
    @wraps(view)
    def wrapped(*args, **kwargs):
        if not session.get("is_admin"):
            return redirect(url_for("admin_login", next=request.path))
        return view(*args, **kwargs)

    return wrapped


def parse_item_rows(item_key, detail_key):
    """Zip two parallel form-array fields into a list of {item, detail} dicts,
    dropping rows where the item name is blank."""
    items = request.form.getlist(item_key)
    details = request.form.getlist(detail_key)
    rows = []
    for i, item in enumerate(items):
        item = item.strip()
        if not item:
            continue
        detail = details[i].strip() if i < len(details) else ""
        rows.append({"item": item, "detail": detail})
    return rows


@app.route("/")
def index():
    events = Event.query.filter_by(active=True).order_by(Event.event_date.desc()).all()
    return render_template("event_select.html", events=events)


@app.route("/event/<int:event_id>")
def event_form(event_id):
    event = Event.query.filter_by(id=event_id, active=True).first()
    if not event:
        flash("That event isn't open for reporting. Please pick again.", "error")
        return redirect(url_for("index"))
    return render_template("event_form.html", event=event)


@app.route("/event/<int:event_id>/qr.svg")
def event_qr(event_id):
    event = Event.query.filter_by(id=event_id, active=True).first_or_404()
    form_url = url_for("event_form", event_id=event.id, _external=True)
    img = qrcode.make(form_url, image_factory=qrcode.image.svg.SvgPathImage)
    buf = io.BytesIO()
    img.save(buf)
    return Response(buf.getvalue(), mimetype="image/svg+xml")


@app.route("/submit", methods=["POST"])
def submit():
    event_id = request.form.get("event_id")
    tech_name = request.form.get("tech_name", "").strip()

    if not event_id or not tech_name:
        flash("Please select your event and enter your name.", "error")
        return redirect(url_for("index"))

    event = Event.query.get(event_id)
    if not event:
        flash("That event could not be found. Please pick again.", "error")
        return redirect(url_for("index"))

    misplaced = parse_item_rows("misplaced_item", "misplaced_kit")
    missing = parse_item_rows("missing_item", "missing_notes")
    broken = parse_item_rows("broken_item", "broken_notes")
    notes = request.form.get("notes", "").strip()

    submission = Submission(
        event_id=event.id,
        event_name=event.name,
        event_date=event.event_date,
        tech_name=tech_name,
        misplaced_items=json.dumps(misplaced),
        missing_items=json.dumps(missing),
        broken_items=json.dumps(broken),
        notes=notes,
    )
    event_name = event.name
    event_date = event.event_date

    db.session.add(submission)
    db.session.commit()

    return render_template("thanks.html", event_name=event_name, event_date=event_date, tech_name=tech_name)


@app.route("/admin/login", methods=["GET", "POST"])
def admin_login():
    if request.method == "POST":
        pin = request.form.get("pin", "")
        if pin == ADMIN_PIN:
            session["is_admin"] = True
            next_url = request.form.get("next") or url_for("admin_events")
            return redirect(next_url)
        flash("Incorrect PIN.", "error")
    return render_template("admin_login.html", next=request.args.get("next", ""))


@app.route("/admin/logout")
def admin_logout():
    session.pop("is_admin", None)
    return redirect(url_for("admin_login"))


def all_venues():
    rows = (
        db.session.query(Event.venue)
        .filter(Event.venue.isnot(None), Event.venue != "")
        .distinct()
        .order_by(Event.venue.asc())
        .all()
    )
    return [r[0] for r in rows]


@app.route("/admin/events", methods=["GET", "POST"])
@admin_required
def admin_events():
    if request.method == "POST":
        name = request.form.get("name", "").strip()
        event_date_raw = request.form.get("event_date", "").strip()
        venue = request.form.get("venue", "").strip()
        if not name or not event_date_raw:
            flash("Event name and date are required.", "error")
        else:
            try:
                event_date = datetime.strptime(event_date_raw, "%Y-%m-%d").date()
            except ValueError:
                flash("Invalid date.", "error")
                return redirect(url_for("admin_events"))
            db.session.add(Event(name=name, event_date=event_date, venue=venue or None))
            db.session.commit()
            flash(f'Added "{name}".', "success")
        return redirect(url_for("admin_events"))

    events = Event.query.order_by(Event.event_date.desc()).all()
    return render_template(
        "admin_events.html",
        events=events,
        today=date.today().isoformat(),
        all_venues=all_venues(),
        pending_count=len(pending_events_query().all()),
    )


@app.route("/admin/events/<int:event_id>/venue", methods=["POST"])
@admin_required
def set_event_venue(event_id):
    event = Event.query.get_or_404(event_id)
    venue = request.form.get("venue", "").strip()
    event.venue = venue or None
    db.session.commit()
    flash(f'Venue {"set" if venue else "cleared"} for "{event.name}".', "success")
    return redirect(url_for("admin_events"))


@app.route("/admin/events/<int:event_id>/date", methods=["POST"])
@admin_required
def set_event_date(event_id):
    event = Event.query.get_or_404(event_id)
    date_raw = request.form.get("event_date", "").strip()
    try:
        new_date = datetime.strptime(date_raw, "%Y-%m-%d").date()
    except ValueError:
        flash("Invalid date.", "error")
        return redirect(url_for("admin_events"))

    event.event_date = new_date
    # Keep already-filed submissions' snapshotted date in sync too.
    Submission.query.filter_by(event_id=event.id).update({"event_date": new_date})
    db.session.commit()
    flash(f'Updated date for "{event.name}" to {new_date.strftime("%-m/%-d/%Y")}.', "success")
    return redirect(url_for("admin_events"))


@app.route("/admin/events/<int:event_id>/toggle", methods=["POST"])
@admin_required
def toggle_event(event_id):
    event = Event.query.get_or_404(event_id)
    event.active = not event.active
    db.session.commit()
    return redirect(url_for("admin_events"))


@app.route("/admin/events/<int:event_id>/delete", methods=["POST"])
@admin_required
def delete_event(event_id):
    event = Event.query.get_or_404(event_id)
    db.session.delete(event)
    db.session.commit()
    flash(f'Deleted "{event.name}".', "success")
    return redirect(url_for("admin_events"))


def build_report(days):
    since = datetime.utcnow() - timedelta(days=days)
    submissions = (
        Submission.query.filter(Submission.submitted_at >= since)
        .order_by(Submission.event_date.desc(), Submission.submitted_at.desc())
        .all()
    )
    report_rows = []
    for s in submissions:
        report_rows.append(
            {
                "id": s.id,
                "event_name": s.event_name,
                "event_date": s.event_date.isoformat(),
                "venue": s.event.venue if s.event else None,
                "tech_name": s.tech_name,
                "misplaced_items": json.loads(s.misplaced_items),
                "missing_items": json.loads(s.missing_items),
                "broken_items": json.loads(s.broken_items),
                "notes": s.notes,
                "submitted_at": s.submitted_at.isoformat(),
            }
        )
    return report_rows


@app.route("/admin/report")
@admin_required
def admin_report():
    days = request.args.get("days", 7, type=int)
    rows = build_report(days)
    all_events = Event.query.order_by(Event.event_date.desc()).all()
    return render_template("report.html", rows=rows, days=days, all_events=all_events)


@app.route("/admin/submissions/<int:submission_id>/reassign", methods=["POST"])
@admin_required
def reassign_submission(submission_id):
    submission = Submission.query.get_or_404(submission_id)
    target_event_id = request.form.get("event_id", type=int)
    target_event = Event.query.get(target_event_id) if target_event_id else None
    if not target_event:
        flash("Pick an event to reassign this submission to.", "error")
        return redirect(url_for("admin_report", days=request.form.get("days", 7, type=int)))

    old_label = f"{submission.event_name} ({submission.event_date.strftime('%-m/%-d/%Y')})"
    submission.event_id = target_event.id
    submission.event_name = target_event.name
    submission.event_date = target_event.event_date
    db.session.commit()
    flash(
        f'Moved {submission.tech_name}\'s submission from "{old_label}" to "{target_event.name}".',
        "success",
    )
    return redirect(url_for("admin_report", days=request.form.get("days", 7, type=int)))


@app.route("/admin/submissions/<int:submission_id>/delete", methods=["POST"])
@admin_required
def delete_submission(submission_id):
    submission = Submission.query.get_or_404(submission_id)
    label = f"{submission.tech_name}'s submission for \"{submission.event_name}\""
    db.session.delete(submission)
    db.session.commit()
    flash(f"Deleted {label}.", "success")
    return redirect(url_for("admin_report", days=request.form.get("days", 7, type=int)))


def pending_events_query():
    """Past events with no submission filed against them yet."""
    reported_event_ids = {
        row[0]
        for row in db.session.query(Submission.event_id).filter(Submission.event_id.isnot(None)).distinct().all()
    }
    return (
        Event.query.filter(Event.event_date <= date.today(), ~Event.id.in_(reported_event_ids))
        .order_by(Event.event_date.desc())
    )


@app.route("/admin/events/pending")
@admin_required
def pending_events():
    events = pending_events_query().all()

    weeks = []
    weeks_by_friday = {}
    for event in events:
        friday = _friday_for_date(event.event_date)
        if friday not in weeks_by_friday:
            weeks_by_friday[friday] = {"friday": friday, "sunday": friday + timedelta(days=2), "events": []}
            weeks.append(weeks_by_friday[friday])
        weeks_by_friday[friday]["events"].append(event)
    weeks.sort(key=lambda w: w["friday"], reverse=True)

    return render_template("pending_events.html", weeks=weeks, pending_count=len(events))


def _friday_for_date(d):
    """The Friday that starts the Fri-Sun weekend containing (or, for a
    Mon-Thu date, most recently preceding) `d`."""
    weekday = d.weekday()  # Monday=0 ... Sunday=6
    if weekday >= 4:
        return d - timedelta(days=weekday - 4)
    return d - timedelta(days=weekday + 3)


def current_weekend_bounds():
    """(friday, sunday) for the weekend most relevant to today: the current
    weekend if today is Fri/Sat/Sun, otherwise the last one."""
    friday = _friday_for_date(date.today())
    return friday, friday + timedelta(days=2)


def _item_list(items):
    """Compile possibly-repeated item dicts (from multiple submissions on the
    same event) into one readable, deduped list: 'Wireless Mic (x2) — left
    at venue, Truss Clamp'."""
    counts = {}
    details = {}
    order = []
    for i in items:
        name = i["item"]
        if name not in counts:
            counts[name] = 0
            details[name] = []
            order.append(name)
        counts[name] += 1
        detail = (i.get("detail") or "").strip()
        if detail and detail not in details[name]:
            details[name].append(detail)

    parts = []
    for name in order:
        label = name if counts[name] == 1 else f"{name} (x{counts[name]})"
        if details[name]:
            label += " — " + "; ".join(details[name])
        parts.append(label)
    return parts


def build_weekend_report(friday=None, sunday=None):
    """Synthesize the given weekend's submissions into a narrative summary,
    with all submissions for the same event compiled into a single entry,
    rather than one row per person who submitted. Defaults to the current
    (or most recently completed) weekend."""
    if friday is None or sunday is None:
        friday, sunday = current_weekend_bounds()
    submissions = (
        Submission.query.filter(Submission.event_date >= friday, Submission.event_date <= sunday)
        .order_by(Submission.event_date.asc(), Submission.event_name.asc(), Submission.submitted_at.asc())
        .all()
    )

    events = []
    events_by_key = {}
    totals = {"missing": 0, "broken": 0, "misplaced": 0}

    for s in submissions:
        key = (s.event_name, s.event_date)
        if key not in events_by_key:
            events_by_key[key] = {
                "event_name": s.event_name,
                "event_date": s.event_date,
                "venue": s.event.venue if s.event else None,
                "missing": [],
                "broken": [],
                "misplaced": [],
                "notes": [],
            }
            events.append(events_by_key[key])
        ev = events_by_key[key]

        m_items = json.loads(s.missing_items)
        b_items = json.loads(s.broken_items)
        p_items = json.loads(s.misplaced_items)
        ev["missing"].extend(m_items)
        ev["broken"].extend(b_items)
        ev["misplaced"].extend(p_items)
        totals["missing"] += len(m_items)
        totals["broken"] += len(b_items)
        totals["misplaced"] += len(p_items)
        if s.notes and s.notes not in ev["notes"]:
            ev["notes"].append(s.notes)

    highlights = []
    clean_events = 0
    event_issue_counts = {}

    for ev in events:
        issue_count = len(ev["missing"]) + len(ev["broken"]) + len(ev["misplaced"])
        if issue_count == 0 and not ev["notes"]:
            clean_events += 1
            continue

        lines = []
        if ev["missing"]:
            lines.append("Missing: " + ", ".join(_item_list(ev["missing"])))
        if ev["broken"]:
            lines.append("Broken/damaged: " + ", ".join(_item_list(ev["broken"])))
        if ev["misplaced"]:
            lines.append("Misplaced: " + ", ".join(_item_list(ev["misplaced"])))

        highlights.append(
            {
                "event_name": ev["event_name"],
                "event_date": ev["event_date"],
                "venue": ev["venue"],
                "lines": lines,
                "notes": ev["notes"],
            }
        )
        event_issue_counts[ev["event_name"]] = issue_count

    stats = {
        "events": len(events),
        "clean_events": clean_events,
        "missing": totals["missing"],
        "broken": totals["broken"],
        "misplaced": totals["misplaced"],
    }

    narrative = build_weekend_narrative(friday, sunday, stats, event_issue_counts)

    return friday, sunday, stats, narrative, highlights


def build_weekend_narrative(friday, sunday, stats, event_issue_counts):
    date_range = f"{friday.strftime('%B %-d')}–{sunday.strftime('%-d, %Y')}"

    if stats["events"] == 0:
        return f"No return reports were submitted for the weekend of {date_range}."

    total_flagged = stats["missing"] + stats["broken"] + stats["misplaced"]

    if total_flagged == 0:
        tone = "a clean weekend across the board"
    elif total_flagged <= 2:
        tone = "a mostly clean weekend, with just a couple of items to follow up on"
    elif total_flagged <= 6:
        tone = "a fairly typical weekend, with a handful of items that need attention"
    else:
        tone = "a rougher weekend than usual, with a notable number of items flagged"

    sentence = f"Across {stats['events']} event{'s' if stats['events'] != 1 else ''} over {date_range}, it was {tone}."

    if total_flagged:
        breakdown = []
        if stats["missing"]:
            breakdown.append(f"{stats['missing']} missing")
        if stats["broken"]:
            breakdown.append(f"{stats['broken']} broken or damaged")
        if stats["misplaced"]:
            breakdown.append(f"{stats['misplaced']} misplaced")
        sentence += (
            f" In total, {', '.join(breakdown)} item{'s' if total_flagged != 1 else ''} "
            f"{'were' if total_flagged != 1 else 'was'} flagged."
        )

    if event_issue_counts:
        worst_event = max(event_issue_counts, key=event_issue_counts.get)
        if len(event_issue_counts) > 1 or event_issue_counts[worst_event] > 1:
            sentence += f" {worst_event} accounted for the most flagged items."

    if stats["clean_events"]:
        sentence += (
            f" {stats['clean_events']} of {stats['events']} event{'s' if stats['events'] != 1 else ''} "
            f"reported a fully clean return."
        )

    return sentence


def build_weekend_docx(friday, sunday, stats, narrative, highlights):
    doc = Document()

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = BRAND_INK

    logo_p = doc.add_paragraph()
    logo_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    logo_p.add_run().add_picture(LOGO_PATH, width=Inches(2.2))

    title = doc.add_heading("Weekend Return Report", level=1)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in title.runs:
        run.font.color.rgb = BRAND_RED

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_run = subtitle.add_run(f"{friday.strftime('%A, %B %-d')} – {sunday.strftime('%A, %B %-d, %Y')}")
    sub_run.italic = True
    sub_run.font.color.rgb = BRAND_GRAY

    doc.add_paragraph()

    doc.add_heading("Executive Summary", level=2)
    doc.add_paragraph(narrative)

    table = doc.add_table(rows=2, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    stat_values = [
        (str(stats["events"]), "Events"),
        (str(stats["clean_events"]), "Clean events"),
        (str(stats["missing"] + stats["broken"] + stats["misplaced"]), "Items flagged"),
    ]
    for col, (num, label) in enumerate(stat_values):
        num_p = table.cell(0, col).paragraphs[0]
        num_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        num_run = num_p.add_run(num)
        num_run.bold = True
        num_run.font.size = Pt(22)
        num_run.font.color.rgb = BRAND_RED

        label_p = table.cell(1, col).paragraphs[0]
        label_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        label_run = label_p.add_run(label)
        label_run.font.size = Pt(9)
        label_run.font.color.rgb = BRAND_GRAY

    doc.add_paragraph()

    doc.add_heading("Highlights", level=2)
    if highlights:
        for h in highlights:
            event_p = doc.add_paragraph()
            venue_suffix = f" — {h['venue']}" if h.get("venue") else ""
            event_run = event_p.add_run(f"{h['event_name']} ({h['event_date'].strftime('%-m/%-d/%Y')}){venue_suffix}")
            event_run.bold = True
            event_run.font.size = Pt(13)

            for line in h["lines"]:
                doc.add_paragraph(line, style="List Bullet")

            if not h["lines"] and h["notes"]:
                no_items_p = doc.add_paragraph()
                no_items_run = no_items_p.add_run("No items flagged — feedback below.")
                no_items_run.font.color.rgb = BRAND_GRAY
                no_items_run.font.size = Pt(9)

            for note in h["notes"]:
                note_p = doc.add_paragraph()
                note_run = note_p.add_run(f"“{note}”")
                note_run.italic = True
                note_run.font.color.rgb = BRAND_GRAY
    else:
        doc.add_paragraph("Nothing flagged this weekend — every event reported a clean return.")

    footer_p = doc.sections[0].footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer_p.add_run(
        f"Siagel Productions · Post-Event Report · Generated {datetime.utcnow().strftime('%-m/%-d/%Y %-I:%M %p')} UTC"
    )
    footer_run.font.size = Pt(8)
    footer_run.font.color.rgb = BRAND_GRAY

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf


def _requested_weekend():
    """Parse ?friday=YYYY-MM-DD from the query string; falls back to the
    current weekend for a missing or invalid value."""
    raw = request.args.get("friday", "")
    if raw:
        try:
            friday = datetime.strptime(raw, "%Y-%m-%d").date()
            return friday, friday + timedelta(days=2)
        except ValueError:
            pass
    return current_weekend_bounds()


@app.route("/admin/report/weekend")
@admin_required
def admin_report_weekend():
    friday, sunday = _requested_weekend()
    friday, sunday, stats, narrative, highlights = build_weekend_report(friday, sunday)
    current_friday, _ = current_weekend_bounds()
    return render_template(
        "report_weekend.html",
        friday=friday,
        sunday=sunday,
        stats=stats,
        narrative=narrative,
        highlights=highlights,
        prev_friday=friday - timedelta(days=7),
        next_friday=friday + timedelta(days=7),
        is_current=(friday == current_friday),
        is_future=(friday > current_friday),
    )


@app.route("/admin/report/weekend/archive")
@admin_required
def admin_report_weekend_archive():
    event_dates = [row[0] for row in db.session.query(Submission.event_date).distinct().all()]
    fridays = sorted({_friday_for_date(d) for d in event_dates}, reverse=True)

    weeks = []
    for friday in fridays:
        sunday = friday + timedelta(days=2)
        _, _, stats, _, _ = build_weekend_report(friday, sunday)
        weeks.append({"friday": friday, "sunday": sunday, "stats": stats})

    return render_template("report_weekend_archive.html", weeks=weeks)


@app.route("/admin/report/weekend/download")
@admin_required
def admin_report_weekend_download():
    friday, sunday = _requested_weekend()
    friday, sunday, stats, narrative, highlights = build_weekend_report(friday, sunday)
    buf = build_weekend_docx(friday, sunday, stats, narrative, highlights)
    filename = f"weekend-report-{friday.isoformat()}-to-{sunday.isoformat()}.docx"
    return send_file(
        buf,
        mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        as_attachment=True,
        download_name=filename,
    )


VENUE_KEYWORDS = {
    "Load-in/out": [
        "load-in", "load in", "load-out", "load out", "loading dock", "dock",
        "parking", "stairs", "elevator", "freight", "ramp",
    ],
    "Power": ["power", "outlet", "circuit", "breaker", "generator", "electrical"],
    "Staff & coordination": [
        "coordinator", "venue staff", "venue manager", "security", "point of contact",
    ],
    "Restrictions": [
        "curfew", "haze", "fog ban", "no rigging", "rigging", "noise ordinance",
        "permit", "fire marshal",
    ],
    "Space & amenities": [
        "ballroom", "green room", "backstage", "wifi", "wi-fi", "cell service", "neighbor",
    ],
}


def categorize_venue_note(text_value):
    """Return the venue-feedback categories a free-text note touches on, by
    keyword match. An empty list means the note doesn't look venue-related."""
    lowered = text_value.lower()
    return [category for category, keywords in VENUE_KEYWORDS.items() if any(k in lowered for k in keywords)]


def build_venue_profiles():
    """One profile per venue, built from any submission notes/item details
    that mention it, grouped by event. Venues with events but no matching
    feedback still show up, with an empty mention list."""
    events = Event.query.filter(Event.venue.isnot(None), Event.venue != "").all()
    events_by_venue = {}
    for e in events:
        events_by_venue.setdefault(e.venue, []).append(e)

    profiles = []
    for venue_name, venue_events in events_by_venue.items():
        event_ids = [e.id for e in venue_events]
        submissions = (
            Submission.query.filter(Submission.event_id.in_(event_ids))
            .order_by(Submission.event_date.desc(), Submission.submitted_at.desc())
            .all()
        )

        mentions = []
        for s in submissions:
            candidates = []
            if s.notes:
                candidates.append(s.notes.strip())
            for key in ("missing_items", "broken_items", "misplaced_items"):
                for item in json.loads(getattr(s, key)):
                    detail = (item.get("detail") or "").strip()
                    if detail:
                        candidates.append(f"{item['item']}: {detail}")

            for text_value in candidates:
                categories = categorize_venue_note(text_value)
                if categories:
                    mentions.append(
                        {
                            "text": text_value,
                            "categories": categories,
                            "event_name": s.event_name,
                            "event_date": s.event_date,
                            "tech_name": s.tech_name,
                        }
                    )

        category_counts = {}
        for m in mentions:
            for c in m["categories"]:
                category_counts[c] = category_counts.get(c, 0) + 1

        profiles.append(
            {
                "venue": venue_name,
                "event_count": len(venue_events),
                "submission_count": len(submissions),
                "mentions": mentions,
                "category_counts": category_counts,
            }
        )

    profiles.sort(key=lambda p: len(p["mentions"]), reverse=True)
    return profiles


@app.route("/admin/venues")
@admin_required
def venues():
    return render_template("venues.html", profiles=build_venue_profiles())


@app.route("/admin/venues/<path:venue_name>")
@admin_required
def venue_detail(venue_name):
    profile = next((p for p in build_venue_profiles() if p["venue"] == venue_name), None)
    if not profile:
        abort(404)
    return render_template("venue_detail.html", profile=profile)


@app.route("/api/report")
def api_report():
    token = request.args.get("token", "")
    if not REPORT_TOKEN or token != REPORT_TOKEN:
        abort(403)
    days = request.args.get("days", 7, type=int)
    rows = build_report(days)
    return jsonify(rows=rows, days=days, generated_at=datetime.utcnow().isoformat())


@app.route("/report/email")
def report_email():
    token = request.args.get("token", "")
    if not REPORT_TOKEN or token != REPORT_TOKEN:
        abort(403)
    days = request.args.get("days", 7, type=int)
    rows = build_report(days)
    return render_template("report_email.html", rows=rows, days=days)


if __name__ == "__main__":
    app.run(debug=True, host="0.0.0.0", port=5050)
